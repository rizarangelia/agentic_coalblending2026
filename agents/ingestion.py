"""Agent 1 — Data Ingestion & Validator.
Deterministic core: parse CSV/XLSX -> standard JSON. LLM (persona prompts/01) hanya
menarasikan alert, TIDAK PERNAH mengubah angka.
"""
import csv, io, json

CV_RANGE = (3000, 6500)

def _read_rows(file_bytes: bytes, filename: str):
    if filename.lower().endswith(('.xlsx', '.xls')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        header = [str(h).strip() for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
    text = file_bytes.decode('utf-8-sig')
    return list(csv.DictReader(io.StringIO(text)))

def ingest(files: dict, manifest: dict):
    """files: {'opening_stock': (bytes, fname), 'coal_getting':..., 'demand_deal':..., 'plan_kapasitas':...}
    Return (data, alerts). data=None jika REJECTED."""
    alerts, data = [], {}
    sites, brands = set(manifest['sites']), set(manifest['brands'])
    specs = {
        'opening_stock':  ('site', ['volume_mt', 'cv_gar']),
        'coal_getting':   ('site', ['volume_mt', 'cv_gar']),
        'demand_deal':    ('brand', ['volume_deal_mt']),
        'plan_kapasitas': ('jetty', ['kapasitas_mt']),
    }
    for key, (idcol, numcols) in specs.items():
        if key not in files:
            alerts.append(('CRITICAL', 'INGEST', f'File {key} belum diunggah — pipeline berhenti (cadence KB-12).'))
            continue
        try:
            rows = _read_rows(*files[key])
        except Exception as e:
            alerts.append(('CRITICAL', 'INGEST', f'{key}: gagal parsing ({e}).'))
            continue
        parsed = {}
        for i, row in enumerate(rows, start=2):
            name = str(row.get(idcol, '')).strip()
            ref = sites if idcol in ('site', 'jetty') else brands
            if name not in ref:
                alerts.append(('CRITICAL', 'INGEST', f'{key} baris {i}: "{name}" tidak dikenal di manifest ({idcol}).'))
                continue
            rec = {}
            for col in numcols:
                v = row.get(col)
                try:
                    v = float(v)
                except (TypeError, ValueError):
                    alerts.append(('CRITICAL', 'INGEST', f'{key} baris {i} ({name}): {col} kosong/non-numerik → REJECTED.'))
                    v = None
                if v is not None and v < 0:
                    alerts.append(('CRITICAL', 'INGEST', f'{key} baris {i} ({name}): {col} negatif ({v}) → REJECTED.'))
                    v = None
                if v is not None and col == 'cv_gar' and not (CV_RANGE[0] <= v <= CV_RANGE[1]):
                    alerts.append(('WARNING', 'INGEST', f'{key} ({name}): CV {v:.0f} di luar rentang wajar {CV_RANGE} — cek satuan GAR.'))
                rec[col] = v
            if all(rec.get(c) is not None for c in numcols):
                parsed[name] = rec
        data[key] = parsed
    # kelengkapan entitas
    for key, ref, idn in (('opening_stock', sites, 'site'), ('coal_getting', sites, 'site'),
                          ('demand_deal', brands, 'brand'), ('plan_kapasitas', sites, 'jetty')):
        missing = ref - set(data.get(key, {}))
        if missing:
            alerts.append(('CRITICAL', 'INGEST', f'{key}: {idn} tidak lengkap, hilang: {sorted(missing)}.'))
    status = 'REJECTED' if any(a[0] == 'CRITICAL' for a in alerts) else 'VALID'
    return ({'status': status, **data} if status == 'VALID' else None), alerts
