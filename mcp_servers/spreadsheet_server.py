"""MCP server — CoalBlend Data (spreadsheet ingestion).

Membungkus agents/ingestion.py (parser + validator CSV/XLSX yang sudah ada) sebagai
tools MCP, supaya Claude Desktop bisa membaca & memvalidasi 4 file input
(opening_stock, coal_getting, demand_deal, plan_kapasitas) tanpa duplikasi logic.

Jalankan manual untuk tes cepat:
    python mcp_servers/spreadsheet_server.py
Daftarkan di Claude Desktop lewat claude_desktop_config.json (lihat mcp_servers/README.md).
"""
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from agents.ingestion import ingest, _read_rows  # noqa: E402

from mcp.server.mcpserver import MCPServer  # noqa: E402

mcp = MCPServer("coalblend-data")

SPREADSHEET_EXT = ('.csv', '.xlsx', '.xls')


def _resolve(path: str) -> Path:
    p = Path(path)
    if not p.is_absolute():
        p = (ROOT / path).resolve()
    return p


def _load_manifest() -> dict:
    return json.loads((ROOT / 'config' / 'constraint_manifest.json').read_text(encoding='utf-8'))


@mcp.tool()
def list_data_files(directory: str = "fixtures") -> list[dict]:
    """List file spreadsheet (.csv/.xlsx/.xls) di sebuah folder.

    directory: path relatif terhadap root project, atau path absolut.
        Default "fixtures" (data simulasi bawaan project).
    """
    d = _resolve(directory)
    if not d.is_dir():
        return [{"error": f"Folder tidak ditemukan: {d}"}]
    out = []
    for f in sorted(d.iterdir()):
        if f.is_file() and f.suffix.lower() in SPREADSHEET_EXT:
            stat = f.stat()
            out.append({
                "name": f.name,
                "path": str(f),
                "size_bytes": stat.st_size,
                "modified": stat.st_mtime,
            })
    return out


@mcp.tool()
def read_spreadsheet(path: str, max_rows: int = 50) -> dict:
    """Baca isi satu file spreadsheet (CSV/XLSX) untuk inspeksi cepat.

    path: path relatif (terhadap root project) atau absolut ke file .csv/.xlsx/.xls.
    max_rows: jumlah baris contoh yang dikembalikan (data lengkap tetap ada di file).
    """
    p = _resolve(path)
    if not p.is_file():
        return {"error": f"File tidak ditemukan: {p}"}
    if p.suffix.lower() not in SPREADSHEET_EXT:
        return {"error": f"Ekstensi tidak didukung: {p.suffix}"}
    try:
        rows = _read_rows(p.read_bytes(), p.name)
    except Exception as e:
        return {"error": f"Gagal parsing {p.name}: {e}"}
    columns = list(rows[0].keys()) if rows else []
    return {
        "file": p.name,
        "columns": columns,
        "total_rows": len(rows),
        "sample_rows": rows[:max_rows],
    }


@mcp.tool()
def get_manifest() -> dict:
    """Ambil constraint_manifest.json: daftar site & brand yang valid, plus aturan
    blending. Pakai ini dulu sebelum ingest_coal_data, supaya tahu nama entity yang sah.
    """
    return _load_manifest()


@mcp.tool()
def ingest_coal_data(
    opening_stock_path: str,
    coal_getting_path: str,
    demand_deal_path: str,
    plan_kapasitas_path: str,
) -> dict:
    """Parse + validasi 4 file input coal blending jadi JSON standar siap dipakai
    agent optimizer/blending. Validasi: entity dikenal di manifest, angka numerik,
    non-negatif, CV dalam rentang wajar (3000-6500).

    Tiap argumen adalah path (relatif ke root project, atau absolut) ke file
    .csv/.xlsx/.xls. Kosongkan salah satu path (string kosong) untuk pakai fixtures
    simulasi bawaan pada slot tersebut.
    """
    manifest = _load_manifest()
    fixtures = {
        'opening_stock_path': 'fixtures/opening_stock_simulasi.csv',
        'coal_getting_path': 'fixtures/coal_getting_simulasi.csv',
        'demand_deal_path': 'fixtures/demand_deal_simulasi.csv',
        'plan_kapasitas_path': 'fixtures/plan_kapasitas_simulasi.csv',
    }
    given = {
        'opening_stock': opening_stock_path,
        'coal_getting': coal_getting_path,
        'demand_deal': demand_deal_path,
        'plan_kapasitas': plan_kapasitas_path,
    }
    files = {}
    resolved_paths = {}
    for key, arg_name in (('opening_stock', 'opening_stock_path'), ('coal_getting', 'coal_getting_path'),
                           ('demand_deal', 'demand_deal_path'), ('plan_kapasitas', 'plan_kapasitas_path')):
        raw_path = given[key] or fixtures[arg_name]
        p = _resolve(raw_path)
        if not p.is_file():
            return {"status": "REJECTED", "data": None,
                     "alerts": [{"level": "CRITICAL", "agent": "INGEST", "message": f"File tidak ditemukan untuk {key}: {p}"}]}
        files[key] = (p.read_bytes(), p.name)
        resolved_paths[key] = str(p)

    data, alerts = ingest(files, manifest)
    return {
        "status": (data or {}).get('status', 'REJECTED'),
        "data": data,
        "alerts": [{"level": lvl, "agent": agent, "message": msg} for lvl, agent, msg in alerts],
        "sources": resolved_paths,
    }


def _sheet_to_xlsx_bytes(ws) -> bytes:
    """Salin satu sheet ke workbook baru berdiri sendiri, supaya bisa dilewatkan ke
    ingest() (yang membaca file per-file) tanpa duplikasi logic parsing/validasi."""
    from openpyxl import Workbook
    wb2 = Workbook()
    ws2 = wb2.active
    for row in ws.iter_rows(values_only=True):
        ws2.append(row)
    buf = io.BytesIO()
    wb2.save(buf)
    return buf.getvalue()


@mcp.tool()
def ingest_coal_data_from_workbook(workbook_path: str) -> dict:
    """Parse + validasi 1 workbook Excel dengan 4 sheet (opening_stock, coal_getting,
    demand_deal, plan_kapasitas) — dipakai untuk template konsolidasi
    fixtures/coalblend_input_template.xlsx. Validasi sama persis dengan ingest_coal_data.

    workbook_path: path relatif (ke root project) atau absolut ke file .xlsx.
    """
    from openpyxl import load_workbook

    p = _resolve(workbook_path)
    if not p.is_file():
        return {"status": "REJECTED", "data": None,
                 "alerts": [{"level": "CRITICAL", "agent": "INGEST", "message": f"File tidak ditemukan: {p}"}]}
    required = ['opening_stock', 'coal_getting', 'demand_deal', 'plan_kapasitas']
    wb = load_workbook(p, data_only=True)
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        return {"status": "REJECTED", "data": None,
                 "alerts": [{"level": "CRITICAL", "agent": "INGEST",
                             "message": f"Sheet tidak ditemukan: {missing}. Sheet tersedia: {wb.sheetnames}"}]}

    manifest = _load_manifest()
    files = {s: (_sheet_to_xlsx_bytes(wb[s]), f"{s}.xlsx") for s in required}
    data, alerts = ingest(files, manifest)
    return {
        "status": (data or {}).get('status', 'REJECTED'),
        "data": data,
        "alerts": [{"level": lvl, "agent": agent, "message": msg} for lvl, agent, msg in alerts],
        "source": str(p),
    }


if __name__ == "__main__":
    mcp.run()
